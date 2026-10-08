"""Slide on pass, final review, re-test, deck editing and versioned exports."""

import json
from collections.abc import Iterator
from typing import Any

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from backend.app import TOKEN_HEADER, create_app
from backend.paths import vcpg_home
from backend.schemas import Status
from tests.conftest import FakeBuilder, FakeLLM, wait_for

ANSWER = "Each user's ratings are loaded into a sparse matrix because most films are unrated."
QUOTE = "most films are unrated"


@pytest.fixture
def api() -> Iterator[TestClient]:
    with TestClient(create_app(), base_url="http://127.0.0.1:8000") as c:
        c.headers[TOKEN_HEADER] = c.get("/api/session").json()["token"]
        yield c


def grade(*statuses: Status) -> str:
    items = [
        {"item_index": i, "status": s, "evidence": None if s == "missing" else QUOTE, "reason": "r"}
        for i, s in enumerate(statuses)
    ]
    return json.dumps({"items": items, "gap": "g"})


def project(api: TestClient, pid: int) -> dict[str, Any]:
    return api.get(f"/api/projects/{pid}").json()


def finished_steps(api: TestClient, fake_llm: FakeLLM) -> int:
    """A one-step project taken through build, tour and a two-answer gate."""
    fake_llm.reply_fixture("planner", "plan_one_step.json")
    pid = api.post("/api/projects", json={"brief": "x"}).json()["id"]
    step_id = api.post(f"/api/projects/{pid}/plan/approve").json()["steps"][0]["id"]
    api.post(f"/api/steps/{step_id}/advance", json={"action": "ready"})
    wait_for(lambda: api.get(f"/api/steps/{step_id}").json()["state"] == "reviewed")
    fake_llm.reply("interviewer", "Why a sparse matrix?")
    api.post(f"/api/steps/{step_id}/advance", json={"action": "ready_for_interview"})
    fake_llm.reply("grader", grade("covered", "missing", "missing"))
    fake_llm.reply("interviewer", "And missing ratings?")
    api.post(f"/api/steps/{step_id}/gate/answer", json={"answer": ANSWER})
    fake_llm.reply("grader", grade("covered", "covered", "covered"))
    fake_llm.reply_fixture("slide_writer", "slide_code.json")
    api.post(f"/api/steps/{step_id}/gate/answer", json={"answer": ANSWER})
    return pid


def test_last_step_moves_the_project_to_final_review(
    api: TestClient, fake_llm: FakeLLM, fake_builder: FakeBuilder
) -> None:
    pid = finished_steps(api, fake_llm)
    view = project(api, pid)
    assert view["status"] == "finishing" and view["current_step_id"] is None
    assert view["finish"]["stage"] == "final_review"


def test_final_review_retest_and_deck(
    api: TestClient, fake_llm: FakeLLM, fake_builder: FakeBuilder
) -> None:
    pid = finished_steps(api, fake_llm)
    api.post(f"/api/projects/{pid}/final-review")
    wait_for(lambda: project(api, pid)["finish"]["stage"] == "retest")
    final_task = fake_builder.calls[-1]
    assert final_task["agent"] == "reviewer" and "<mode>\nfinal\n</mode>" in final_task["task"]

    fake_llm.reply("interviewer", "Re-test: how do missing ratings differ from zeros?")
    view = api.post(f"/api/projects/{pid}/retest/start").json()["finish"]["retest"]
    assert (view["passed"], view["total"], view["current_step"]) == (0, 1, "Ratings loader")
    asked = fake_llm.requests_for("interviewer")[-1][0]["content"]
    assert "missing ratings differ from zero" in asked  # the point covered last at the gate

    fake_llm.reply("grader", grade("missing"))
    fake_llm.reply("tutor", "Think of an empty cell as 'never watched'.")
    fake_llm.reply("interviewer", "Try again: what does an empty cell mean?")
    view = api.post(f"/api/projects/{pid}/retest/answer", json={"answer": "no idea"}).json()
    roles = [m["role"] for m in view["finish"]["retest"]["messages"]]
    assert roles == ["assistant", "user", "reteach", "assistant"]

    fake_llm.reply("grader", grade("covered"))
    for kind in ("title", "overview", "retro"):
        fake_llm.reply_fixture("slide_writer", f"slide_{kind}.json")
    view = api.post(f"/api/projects/{pid}/retest/answer", json={"answer": ANSWER}).json()
    assert view["finish"]["stage"] == "deck"
    modes = [r[0]["content"].split("\n")[1] for r in fake_llm.requests_for("slide_writer")]
    assert modes == ["step", "title", "overview", "retrospective"]

    deck = api.get(f"/api/projects/{pid}/deck").json()["slides"]
    assert [s["kind"] for s in deck] == ["title", "overview", "step", "retrospective"]
    assert deck[2]["code"] == ["x = 1"]  # read from the commit, not written by the model

    assert api.post(f"/api/projects/{pid}/deck/done").json()["finish"]["stage"] == "readme"


@pytest.fixture
def deck(api: TestClient, fake_llm: FakeLLM, fake_builder: FakeBuilder) -> tuple[int, list[Any]]:
    pid = finished_steps(api, fake_llm)
    api.post(f"/api/projects/{pid}/final-review")
    wait_for(lambda: project(api, pid)["finish"]["stage"] == "retest")
    fake_llm.reply("interviewer", "q")
    api.post(f"/api/projects/{pid}/retest/start")
    fake_llm.reply("grader", grade("covered"))
    for kind in ("title", "overview", "retro"):
        fake_llm.reply_fixture("slide_writer", f"slide_{kind}.json")
    api.post(f"/api/projects/{pid}/retest/answer", json={"answer": ANSWER})
    return pid, api.get(f"/api/projects/{pid}/deck").json()["slides"]


def test_reorder_edit_and_regenerate(
    api: TestClient, fake_llm: FakeLLM, deck: tuple[int, list[Any]]
) -> None:
    pid, slides = deck
    ids = [s["id"] for s in slides]
    order = [ids[0], ids[2], ids[1], ids[3]]
    view = api.post(f"/api/projects/{pid}/deck/order", json={"slide_ids": order}).json()
    assert [s["id"] for s in view["slides"]] == order
    assert (
        api.post(f"/api/projects/{pid}/deck/order", json={"slide_ids": ids[:2]}).status_code == 409
    )

    edited = dict(slides[0]["slide"], title="My recommender", notes="My own words.")
    assert api.put(f"/api/slides/{ids[0]}", json=edited).status_code == 200
    assert api.put(f"/api/slides/{ids[0]}", json={"title": "broken"}).status_code == 422
    assert (
        api.get(f"/api/projects/{pid}/deck").json()["slides"][0]["slide"]["title"]
        == "My recommender"
    )

    fake_llm.reply_fixture("slide_writer", "slide_step.json")
    api.post(f"/api/slides/{ids[2]}/regenerate", json={"instruction": "use a flow diagram"})
    sent = fake_llm.requests_for("slide_writer")[-1][0]["content"]
    assert "<instruction>\nuse a flow diagram\n</instruction>" in sent
    slide = next(
        s for s in api.get(f"/api/projects/{pid}/deck").json()["slides"] if s["id"] == ids[2]
    )
    assert slide["slide"]["visual"]["kind"] == "flow"


def test_exports_are_versioned_and_complete(api: TestClient, deck: tuple[int, list[Any]]) -> None:
    pid, slides = deck
    for fmt in ("pptx", "pptx", "pdf"):
        assert api.post(f"/api/projects/{pid}/deck/export", json={"format": fmt}).status_code == 200
    exports = api.get(f"/api/projects/{pid}/deck").json()["exports"]
    assert sorted((e["format"], e["version"]) for e in exports) == [
        ("pdf", 1), ("pptx", 1), ("pptx", 2),
    ]  # fmt: skip
    folder = vcpg_home() / "exports" / str(pid)
    assert {p.name for p in folder.iterdir()} == {"deck_v1.pptx", "deck_v2.pptx", "deck_v1.pdf"}

    prs = Presentation(str(folder / "deck_v1.pptx"))
    assert len(prs.slides) == len(slides)
    for page, entry in zip(prs.slides, slides, strict=True):
        assert any(shape.shape_type == 13 for shape in page.shapes)  # the visual, as a picture
        assert page.notes_slide.notes_text_frame.text == entry["slide"]["notes"]
        assert entry["slide"]["title"] in " ".join(
            s.text_frame.text for s in page.shapes if s.has_text_frame
        )

    pdf = (folder / "deck_v1.pdf").read_bytes()
    assert pdf.count(b"/Type /Page\n") + pdf.count(b"/Type /Page ") + pdf.count(
        b"/Type /Page>"
    ) >= len(slides)

    download = api.get(f"/api/exports/{exports[-1]['id']}/download")
    assert download.status_code == 200 and download.content[:2] == b"PK"
