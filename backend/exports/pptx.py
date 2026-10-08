"""PowerPoint export: 16:9 slides from slide JSON, visuals as images, speaker notes kept."""

import io
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from backend.exports.render import visual_figure

BG, TEXT, MUTE, ACC, ACC3 = "0D0A08", "F7EBDD", "BFA58F", "FF9B3D", "FFD27A"
DISPLAY, BODY = "Oxanium", "IBM Plex Sans"  # PowerPoint falls back to Arial / Calibri


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor(*bytes.fromhex(hex_color))


def _text(slide: Any, box: tuple[float, float, float, float], paragraphs: list[tuple[str, int, str, str, bool]]) -> None:
    left, top, width, height = (Inches(v) for v in box)
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    for i, (text, size, color, font, bold) in enumerate(paragraphs):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = para.add_run()
        run.text = text
        run.font.size, run.font.name, run.font.bold = Pt(size), font, bold
        run.font.color.rgb = _rgb(color)


def _png(visual: dict[str, Any], code: list[str] | None) -> io.BytesIO:
    buffer = io.BytesIO()
    visual_figure(visual, code).savefig(buffer, format="png", dpi=160, facecolor="#0D0A08")
    buffer.seek(0)
    return buffer


def export(deck: list[dict[str, Any]], path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for entry in deck:
        data, page = entry["slide"], prs.slides.add_slide(blank)
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = _rgb(BG)
        _text(page, (0.5, 0.3, 12.3, 1.4), [
            (data["kicker"].upper(), 12, ACC, DISPLAY, True),
            (data["title"], 30, TEXT, DISPLAY, True),
            (data["what"], 14, MUTE, BODY, False),
        ])  # fmt: skip
        page.shapes.add_picture(
            _png(data["visual"], entry["code"]), Inches(0.5), Inches(1.9), width=Inches(8.2)
        )
        right: list[tuple[str, int, str, str, bool]] = []
        if data["decisions"]:
            right.append(("DECISIONS", 12, ACC3, DISPLAY, True))
            right += [(f"{d['choice']}: {d['tradeoff']}", 12, TEXT, BODY, False) for d in data["decisions"]]
        if data["failure_and_scale"]:
            right.append(("FAILURE AND SCALE", 12, ACC3, DISPLAY, True))
            right += [(line, 12, TEXT, BODY, False) for line in data["failure_and_scale"]]
        if right:
            _text(page, (9.0, 1.9, 3.9, 4.6), right)
        if data["struggled"]:
            footer = "  ·  ".join(f"{s['question']} ({s['hints']} hints)" for s in data["struggled"])
            _text(page, (0.5, 6.6, 12.3, 0.7), [(f"Questions I struggled with: {footer}", 11, ACC3, BODY, False)])
        page.notes_slide.notes_text_frame.text = data["notes"]
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


